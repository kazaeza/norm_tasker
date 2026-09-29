"""Обезличенный xlsx с той же структурой, что у реального КП.

Все темы, ссылки и имена выдуманы. Комментарии кладутся в файл так же, как это делает
экспорт Google Sheets: частями xl/threadedComments и xl/persons.
"""

from __future__ import annotations

import posixpath
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from io import BytesIO
from xml.etree import ElementTree as ET

import openpyxl

from norm_tasker.kp.comments import (
    NS_MAIN,
    NS_REL_DOC,
    NS_REL_PKG,
    NS_THREADED,
    REL_PERSON,
    REL_THREADED,
    resolve_part,
)

WEEKDAYS = ["понедельник", "вторник", "среда", "четверг", "пятница", "суббота", "воскресенье"]
DOC = "https://docs.google.com/document/d/{}/edit?usp=sharing"


@dataclass
class Cell:
    """Один день недельного блока."""

    rubric: str | None = None
    status: str | None = None
    topic: str | None = None
    doc: str | None = None  # id дока, на который ссылается тема
    extra: str | None = None
    extra_doc: str | None = None


@dataclass
class Week:
    start: date  # понедельник
    days: list[Cell | None] = field(default_factory=lambda: [None] * 7)


@dataclass
class Comment:
    sheet: str
    ref: str
    id: str
    author: str
    text: str
    created: str = "2026-09-20T10:00:00.00"
    parent: str | None = None
    done: bool = False


def build_kp(sheets: dict[str, list[Week]], hidden: tuple[str, ...] = ()) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, weeks in sheets.items():
        ws = wb.create_sheet(title)
        for col, name in enumerate(WEEKDAYS, start=1):
            ws.cell(1, col, name)
        for index, week in enumerate(weeks):
            row = 2 + index * 5
            for col in range(1, 8):
                ws.cell(
                    row,
                    col,
                    datetime.combine(week.start + timedelta(days=col - 1), datetime.min.time()),
                )
                cell = week.days[col - 1]
                if cell is None:
                    continue
                ws.cell(row + 1, col, cell.rubric)
                ws.cell(row + 2, col, cell.status)
                topic = ws.cell(row + 3, col, cell.topic)
                if cell.doc:
                    topic.hyperlink = DOC.format(cell.doc)
                extra = ws.cell(row + 4, col, cell.extra)
                if cell.extra_doc:
                    extra.hyperlink = DOC.format(cell.extra_doc)
        # Заметки справа от календаря — не посты.
        ws.cell(3, 9, "Ссылки:")
        ws.cell(5, 8, "Рабочая заметка")
        if title in hidden:
            ws.sheet_state = "hidden"
    stats = wb.create_sheet("Статистика")
    stats.cell(1, 1, "Дата")
    stats.cell(2, 1, "Охват")
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _sheet_paths(archive: zipfile.ZipFile) -> dict[str, str]:
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {r.get("Id"): r.get("Target") for r in rels.iter(f"{{{NS_REL_PKG}}}Relationship")}
    result = {}
    for sheet in workbook.iter(f"{{{NS_MAIN}}}sheet"):
        target = targets[sheet.get(f"{{{NS_REL_DOC}}}id")]
        result[sheet.get("name")] = resolve_part("xl", target)
    return result


def add_comments(xlsx: bytes, comments: list[Comment]) -> bytes:
    """Добавляет в файл ветки комментариев в формате экспорта Google Sheets."""
    source = zipfile.ZipFile(BytesIO(xlsx))
    files = {name: source.read(name) for name in source.namelist()}
    paths = _sheet_paths(source)

    people: dict[str, str] = {}
    for comment in comments:
        people.setdefault(comment.author, f"{{person-{len(people) + 1}}}")
    person_xml = ET.Element(f"{{{NS_THREADED}}}personList")
    for author, pid in people.items():
        ET.SubElement(
            person_xml,
            f"{{{NS_THREADED}}}person",
            {"displayName": author, "id": pid, "providerId": "google-sheets"},
        )
    files["xl/persons/person.xml"] = ET.tostring(person_xml, xml_declaration=True, encoding="UTF-8")

    workbook_rels = ET.fromstring(files["xl/_rels/workbook.xml.rels"])
    ET.SubElement(
        workbook_rels,
        f"{{{NS_REL_PKG}}}Relationship",
        {"Id": "rIdPersons", "Type": REL_PERSON, "Target": "persons/person.xml"},
    )
    files["xl/_rels/workbook.xml.rels"] = ET.tostring(workbook_rels, xml_declaration=True)

    by_sheet: dict[str, list[Comment]] = {}
    for comment in comments:
        by_sheet.setdefault(comment.sheet, []).append(comment)
    for number, (sheet, items) in enumerate(by_sheet.items(), start=1):
        root = ET.Element(f"{{{NS_THREADED}}}ThreadedComments")
        for item in items:
            attrs = {
                "ref": item.ref,
                "dT": item.created,
                "personId": people[item.author],
                "id": item.id,
                "done": "1" if item.done else "0",
            }
            if item.parent:
                attrs["parentId"] = item.parent
            node = ET.SubElement(root, f"{{{NS_THREADED}}}threadedComment", attrs)
            ET.SubElement(node, f"{{{NS_THREADED}}}text").text = item.text
        comments_path = f"xl/threadedComments/threadedComment{number}.xml"
        files[comments_path] = ET.tostring(root, xml_declaration=True, encoding="UTF-8")

        sheet_path = paths[sheet]
        rels_path = posixpath.join(
            posixpath.dirname(sheet_path), "_rels", posixpath.basename(sheet_path) + ".rels"
        )
        if rels_path in files:
            rels = ET.fromstring(files[rels_path])
        else:
            rels = ET.Element(f"{{{NS_REL_PKG}}}Relationships")
        ET.SubElement(
            rels,
            f"{{{NS_REL_PKG}}}Relationship",
            {
                "Id": f"rIdTc{number}",
                "Type": REL_THREADED,
                "Target": f"../threadedComments/threadedComment{number}.xml",
            },
        )
        files[rels_path] = ET.tostring(rels, xml_declaration=True)

    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    return out.getvalue()


def standard_kp() -> dict[str, list[Week]]:
    """Сентябрь и октябрь 2026 с общей неделей 28.09–04.10 и разными данными в ней."""
    september = [
        Week(
            date(2026, 9, 14),
            [
                Cell("Продукт", "Выпущено", "[док] Тема про поиск", "docA"),
                Cell("Развлекательный", "Выпущено", "Мем про переписку"),
                None,
                None,
                None,
                None,
                None,
            ],
        ),
        Week(
            date(2026, 9, 28),  # устаревшая версия недели: в октябрьском листе она другая
            [
                Cell("Продукт", "Готово", "Старая тема A"),
                Cell("Развлекательный", None, "Старая тема B"),
                None,
                None,
                None,
                Cell("Продукт", None, "Старый мем на субботу"),
                None,
            ],
        ),
    ]
    october = [
        Week(
            date(2026, 9, 28),
            [
                Cell(
                    "Продукт",
                    "Выпущено",
                    "[док] Новая тема A",
                    "docB",
                    "Пост про второй повод",
                    None,
                ),
                Cell("Развлекательный", "В работе ", "Тема B"),
                Cell("Информационный", None, None),
                Cell("Развлекательный", None, None),
                Cell("Информационный", None, "Тема про пятницу", None, "NB: не забыть про акцию"),
                None,
                Cell(None, None, None, None, "заметка без поста"),
            ],
        ),
        Week(
            date(2026, 10, 5),
            [
                Cell("Развлекательный"),
                Cell("Информационный", None, "Тема на вторник"),
                Cell("Продукт", "Текст на согласовании", "[док] Тема на среду", "docC"),
                Cell("Информационный"),
                Cell("Развлекательный", "Странный статус"),
                None,
                None,
            ],
        ),
    ]
    return {"Сентябрь 2026": september, "Октябрь 2026": october}
