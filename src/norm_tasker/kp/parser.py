"""Разбор календаря КП из xlsx-выгрузки Google Sheets.

Шаблон листа (проверен на реальном файле): по листу на месяц; неделя — блок из пяти
строк, колонки A–G — понедельник–воскресенье:

    строка 0: даты
    строка 1: рубрика
    строка 2: статус (выпадающий список)
    строка 3: тема (со ссылкой на док и пометкой «[док]»)
    строка 4: второй пост или заметка

Блоки ищутся по строке с датами, а не по номерам строк. Неделя на стыке месяцев есть
в двух листах: побеждает лист более позднего месяца, если в его блоке есть данные.
"""

from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from io import BytesIO

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from norm_tasker.kp.comments import RawComment, read_threaded_comments
from norm_tasker.kp.models import KpComment, KpParseResult, KpSlot, SheetInfo, clean_topic
from norm_tasker.tracker.stages import DEFAULT_KP_STATUS_STAGES, Stage, normalize

RU_MONTHS = {
    "январь": 1, "февраль": 2, "март": 3, "апрель": 4, "май": 5, "июнь": 6,
    "июль": 7, "август": 8, "сентябрь": 9, "октябрь": 10, "ноябрь": 11, "декабрь": 12,
}  # fmt: skip
RU_MONTHS_GENITIVE = [
    "января", "февраля", "марта", "апреля", "мая", "июня",
    "июля", "августа", "сентября", "октября", "ноября", "декабря",
]  # fmt: skip
SHEET_TITLE = re.compile(r"^\s*([а-яё]+)\s+(\d{4})\s*$", re.IGNORECASE)
# Пятая строка дня — второй пост или заметка. Заметкой считаем только явно помеченное:
# лучше лишний раз спросить «это пост?», чем молча пропустить пост.
NOTE_MARKER = re.compile(r"^\s*(nb|нб|заметка|примечание|ps|пс)\b", re.IGNORECASE)

WEEK_ROWS = 5
DAY_COLUMNS = 7
MIN_DATES_IN_ROW = 5


def month_of_title(title: str) -> date | None:
    """«Октябрь 2026» → 2026-10-01; для листов не по шаблону (статистика и т. п.) — None."""
    match = SHEET_TITLE.match(title)
    if not match:
        return None
    month = RU_MONTHS.get(match.group(1).lower().replace("ё", "е"))
    return date(int(match.group(2)), month, 1) if month else None


def month_title(month: date) -> str:
    """2026-10-01 → «Октябрь 2026»."""
    name = next(k for k, v in RU_MONTHS.items() if v == month.month)
    return f"{name.capitalize()} {month.year}"


def looks_like_note(text: str) -> bool:
    return bool(NOTE_MARKER.match(text))


@dataclass
class _Day:
    day: date
    col: int
    cells: tuple[str, str, str, str, str]  # даты, рубрика, статус, тема, второй пост
    rubric: str | None
    status: str | None
    topic: str | None
    topic_url: str | None
    extra: str | None
    extra_url: str | None


@dataclass
class _Week:
    sheet: str
    month: date
    start: date
    days: list[_Day] = field(default_factory=list)

    @property
    def has_content(self) -> bool:
        return any(d.status or d.topic or d.extra for d in self.days)


def _text(ws: Worksheet, row: int, col: int) -> str | None:
    value = ws.cell(row, col).value
    if isinstance(value, str):
        return value.strip() or None
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int | float):
        return str(value)
    return None  # даты и прочее в текстовых строках не считаем


def _link(ws: Worksheet, row: int, col: int) -> str | None:
    link = ws.cell(row, col).hyperlink
    return link.target if link is not None and link.target else None


def _as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _scan_sheet(ws: Worksheet, month: date, warnings: list[str]) -> list[_Week]:
    weeks: list[_Week] = []
    for row in range(1, ws.max_row + 1):
        dates = {c: _as_date(ws.cell(row, c).value) for c in range(1, DAY_COLUMNS + 1)}
        dates = {c: d for c, d in dates.items() if d is not None}
        if len(dates) < MIN_DATES_IN_ROW:
            continue
        starts = {d - timedelta(days=c - 1) for c, d in dates.items()}
        if len(starts) != 1:
            warnings.append(f"Лист «{ws.title}», строка {row}: даты идут не по порядку")
            continue
        start = starts.pop()
        week = _Week(sheet=ws.title, month=month, start=start)
        for col, day in sorted(dates.items()):
            week.days.append(
                _Day(
                    day=day,
                    col=col,
                    cells=tuple(ws.cell(row + i, col).coordinate for i in range(WEEK_ROWS)),
                    rubric=_text(ws, row + 1, col),
                    status=_text(ws, row + 2, col),
                    topic=_text(ws, row + 3, col),
                    topic_url=_link(ws, row + 3, col),
                    extra=_text(ws, row + 4, col),
                    extra_url=_link(ws, row + 4, col),
                )
            )
        weeks.append(week)
    return weeks


def _pick_winner(candidates: list[_Week]) -> _Week:
    """Из дублей недели выбирает блок более позднего листа, в котором есть данные."""
    ordered = sorted(candidates, key=lambda w: w.month, reverse=True)
    return next((w for w in ordered if w.has_content), ordered[0])


def _slots_of(day: _Day, sheet: str) -> list[KpSlot]:
    slots: list[KpSlot] = []
    if day.rubric or day.status or day.topic:
        slots.append(
            KpSlot(
                date=day.day,
                slot=1,
                rubric=day.rubric,
                status=day.status,
                topic=clean_topic(day.topic),
                doc_url=day.topic_url,
                sheet=sheet,
                cell=day.cells[3] if day.topic else day.cells[1],
            )
        )
        # Пятая строка — второй пост, если она не помечена как заметка и у дня есть основной пост.
        if day.extra and not looks_like_note(day.extra):
            slots.append(
                KpSlot(
                    date=day.day,
                    slot=2,
                    rubric=None,
                    status=None,
                    topic=clean_topic(day.extra),
                    doc_url=day.extra_url,
                    sheet=sheet,
                    cell=day.cells[4],
                )
            )
    return slots


def _cell_index(weeks: list[_Week]) -> dict[tuple[str, str], tuple[date, int]]:
    index: dict[tuple[str, str], tuple[date, int]] = {}
    for week in weeks:
        for day in week.days:
            for coord in day.cells[:4]:
                index[(week.sheet, coord)] = (day.day, 1)
            index[(week.sheet, day.cells[4])] = (day.day, 2)
    return index


def _link_comments(
    raw: list[RawComment],
    index: dict[tuple[str, str], tuple[date, int]],
    existing: set[tuple[date, int]],
) -> list[KpComment]:
    resolved_roots = {c.id: c.resolved for c in raw if c.parent_id is None}
    comments = []
    for item in raw:
        thread = item.parent_id or item.id
        day, slot = index.get((item.sheet, item.cell), (None, None))
        if day is not None and (day, slot) not in existing:
            day, slot = None, None
        comments.append(
            KpComment(
                id=item.id,
                sheet=item.sheet,
                cell=item.cell,
                author=item.author,
                text=item.text,
                created=item.created,
                thread_id=thread,
                resolved=resolved_roots.get(thread, item.resolved),
                day=day,
                slot=slot,
            )
        )
    return comments


def parse_kp(
    data: bytes,
    *,
    since: date | None = None,
    until: date | None = None,
    status_stages: dict[str, Stage] | None = None,
    with_comments: bool = True,
) -> KpParseResult:
    """Разбирает xlsx-выгрузку КП.

    since/until ограничивают, какие листы читать (по месяцу листа), чтобы не разбирать
    архив за годы. Сами посты по датам не фильтруются.
    """
    result = KpParseResult()
    workbook = openpyxl.load_workbook(BytesIO(data), data_only=True)

    first_month = (since.replace(day=1) - timedelta(days=1)).replace(day=1) if since else None
    last_month = until.replace(day=1) if until else None

    weeks_by_start: dict[date, list[_Week]] = defaultdict(list)
    all_weeks: list[_Week] = []
    sheet_weeks: dict[str, int] = {}
    hidden: dict[str, bool] = {}
    months: dict[str, date] = {}

    for ws in workbook.worksheets:
        month = month_of_title(ws.title)
        if month is None:
            continue
        if (first_month and month < first_month) or (last_month and month > last_month):
            continue
        weeks = _scan_sheet(ws, month, result.warnings)
        if not weeks:
            result.warnings.append(
                f"Лист «{ws.title}»: не нашла строк с датами — шаблон изменился?"
            )
        for week in weeks:
            weeks_by_start[week.start].append(week)
        all_weeks.extend(weeks)
        sheet_weeks[ws.title] = len(weeks)
        hidden[ws.title] = ws.sheet_state != "visible"
        months[ws.title] = month

    winners: list[_Week] = [_pick_winner(c) for _, c in sorted(weeks_by_start.items())]
    used = defaultdict(int)
    mapping = status_stages if status_stages is not None else DEFAULT_KP_STATUS_STAGES

    for week in winners:
        used[week.sheet] += 1
        for day in week.days:
            result.calendar_days[day.day] = week.sheet
            slots = _slots_of(day, week.sheet)
            result.slots.extend(slots)
            first = slots[0] if slots else None
            if first and first.status and normalize(first.status) not in mapping:
                result.unknown_statuses.add(first.status)
    result.slots.sort(key=lambda s: s.key)

    for title, count in sheet_weeks.items():
        result.sheets.append(
            SheetInfo(
                name=title,
                month=months[title],
                hidden=hidden[title],
                weeks=count,
                used_weeks=used[title],
            )
        )
    for status in sorted(result.unknown_statuses):
        result.warnings.append(
            f"Неизвестный статус в КП: «{status}» — проверьте настройку статусов"
        )

    if with_comments:
        raw = [c for c in read_threaded_comments(data) if c.sheet in sheet_weeks]
        result.comments = _link_comments(raw, _cell_index(all_weeks), {s.key for s in result.slots})
    return result
